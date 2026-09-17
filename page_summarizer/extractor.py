"""Извлечение текста судебного акта из файла, вставки или публикации по URL."""

from __future__ import annotations

import logging
import re
import time
from io import BytesIO
from pathlib import Path
from typing import Final

import httpx
from bs4 import BeautifulSoup

try:
    import trafilatura
except ImportError:  # pragma: no cover
    trafilatura = None

from .exceptions import ExtractionError, FetchError

logger = logging.getLogger(__name__)

_RETRYABLE_STATUS: Final = {429, 500, 502, 503, 504}

_ALLOWED_CONTENT_TYPES: Final = (
    "text/html",
    "application/xhtml+xml",
    "text/plain",
    "text/markdown",
    "application/pdf",
)

_TEXT_SUFFIXES: Final = {".txt", ".md"}


class _RetryableHTTPError(Exception):
    """Внутренняя ошибка для повторных попыток."""


def normalize_text(text: str) -> str:
    """Нормализует текст: убирает лишние пробелы и пустые строки."""
    text = text.replace("\r", " ")
    text = re.sub(r"[ \t]+", " ", text)
    lines = [line.strip() for line in text.splitlines()]
    text = "\n".join(line for line in lines if line)
    return text.strip()


def fetch_page_html(
    url: str,
    *,
    timeout: float,
    max_retries: int,
    user_agent: str,
) -> str:
    """Загружает HTML или текстовый контент страницы с повторными попытками."""
    normalized_url = url.strip()

    if not normalized_url.lower().startswith(("http://", "https://")):
        raise FetchError(
            "Поддерживаются только ссылки с префиксами http:// или https://."
        )

    headers = {
        "User-Agent": user_agent,
        "Accept": (
            "text/html,application/xhtml+xml,application/pdf;q=0.9,"
            "text/plain;q=0.8,text/markdown;q=0.7,*/*;q=0.6"
        ),
        "Accept-Language": "ru,en;q=0.8",
    }

    last_error: Exception | None = None

    for attempt in range(1, max_retries + 1):
        try:
            with httpx.Client(
                timeout=timeout,
                follow_redirects=True,
                headers=headers,
            ) as client:
                response = client.get(normalized_url)

            if response.status_code in _RETRYABLE_STATUS:
                raise _RetryableHTTPError(f"HTTP {response.status_code}")

            if response.status_code >= 400:
                raise FetchError(
                    f"Сервер вернул ошибку HTTP {response.status_code} "
                    f"для URL {normalized_url}."
                )

            content_type = response.headers.get("content-type", "").lower()
            if content_type and not any(
                content_type.startswith(item) for item in _ALLOWED_CONTENT_TYPES
            ):
                raise ExtractionError(
                    f"Неподдерживаемый тип контента: {content_type or 'неизвестен'}."
                )

            logger.info(
                "Страница загружена: status=%s bytes=%s content-type=%s",
                response.status_code,
                len(response.content),
                content_type or "неизвестен",
            )
            if content_type.startswith("application/pdf") or normalized_url.lower().endswith(
                ".pdf"
            ):
                return extract_from_bytes(response.content, filename="act.pdf")
            return response.text

        except (
            _RetryableHTTPError,
            httpx.TimeoutException,
            httpx.TransportError,
        ) as exc:
            last_error = exc
            if attempt < max_retries:
                delay = min(2**attempt, 8)
                logger.warning(
                    "Попытка %s/%s для %s не удалась: %s. Повтор через %s сек.",
                    attempt,
                    max_retries,
                    normalized_url,
                    exc,
                    delay,
                )
                time.sleep(delay)
            continue

    raise FetchError(
        f"Не удалось загрузить страницу после {max_retries} попыток."
    ) from last_error


def extract_from_file(path: str | Path, *, min_chars: int = 180) -> str:
    """Читает судебный акт с диска (PDF, DOCX, TXT)."""
    file_path = Path(path)
    if not file_path.is_file():
        raise ExtractionError(f"Файл не найден: {file_path}")
    return extract_from_bytes(
        file_path.read_bytes(),
        filename=file_path.name,
        min_chars=min_chars,
    )


def extract_from_bytes(
    data: bytes,
    *,
    filename: str,
    min_chars: int = 180,
) -> str:
    """Извлекает текст из содержимого PDF, DOCX или TXT."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".doc":
        raise ExtractionError(
            "Формат .doc не поддерживается. Сохраните акт как PDF, DOCX или TXT."
        )
    if suffix == ".pdf":
        text = _extract_pdf(data)
    elif suffix == ".docx":
        text = _extract_docx(data)
    elif suffix in _TEXT_SUFFIXES or not suffix:
        text = data.decode("utf-8", errors="replace")
    else:
        raise ExtractionError(
            f"Неподдерживаемый тип файла: {suffix or filename}. "
            "Нужен PDF, DOCX или TXT."
        )

    text = normalize_text(text)
    if len(text) < min_chars:
        raise ExtractionError(
            "В файле слишком мало текста для анализа. "
            "Если это скан без текстового слоя, нужен PDF с текстом или DOCX."
        )
    logger.info("Текст извлечён из %s: %s символов", filename, len(text))
    return text


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise ExtractionError("Для PDF нужен пакет pypdf.") from exc
    reader = PdfReader(BytesIO(data))
    parts = [(page.extract_text() or "") for page in reader.pages]
    return "\n".join(parts)


def _extract_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise ExtractionError("Для DOCX нужен пакет python-docx.") from exc
    document = Document(BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_main_text(html: str, *, url: str, min_chars: int = 180) -> str:
    """Извлекает основной текст страницы без служебного шума."""
    stripped = html.lstrip()
    looks_like_html = stripped.startswith("<") or "<html" in stripped[:500].lower()
    if not looks_like_html:
        text = normalize_text(html)
        if len(text) < min_chars:
            raise ExtractionError(
                "На странице слишком мало текстового контента для анализа."
            )
        logger.info("Текст публикации без HTML: %s символов", len(text))
        return text

    if trafilatura is not None:
        extracted = trafilatura.extract(
            html,
            url=url,
            include_comments=False,
            include_tables=True,
            include_links=False,
            include_images=False,
            favor_recall=True,
        )
        if extracted:
            text = normalize_text(extracted)
            if len(text) >= min_chars:
                logger.info("Текст извлечён через trafilatura: %s символов", len(text))
                return text
            logger.debug(
                "trafilatura дала слишком короткий текст (%s), пробую BeautifulSoup",
                len(text),
            )

    soup = BeautifulSoup(html, "html.parser")
    for tag_name in (
        "script",
        "style",
        "noscript",
        "iframe",
        "svg",
        "form",
        "nav",
        "header",
        "footer",
    ):
        for tag in soup.find_all(tag_name):
            tag.decompose()

    root = soup.find("main") or soup.find("article") or soup.body or soup
    text = normalize_text(root.get_text(" ", strip=True))

    if len(text) < min_chars:
        raise ExtractionError(
            "На странице слишком мало текстового контента для анализа."
        )

    logger.info("Текст извлечён через BeautifulSoup: %s символов", len(text))
    return text
